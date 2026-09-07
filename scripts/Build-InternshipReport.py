from __future__ import annotations

import io
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import win32com.client
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "00-project" / "RAPPORT_STAGE_DISPLAY_CONTROL.docx"
OUTPUT_DIR = ROOT / "docs" / "04-report"
DOCX_OUTPUT = OUTPUT_DIR / "Rapport_PFE_Display_Control_NEGAGZA_Nizar.docx"
PDF_OUTPUT = OUTPUT_DIR / "Rapport_PFE_Display_Control_NEGAGZA_Nizar.pdf"
FIGURES_DIR = OUTPUT_DIR / "figures"
FIGURE_GENERATOR = ROOT / "scripts" / "Generate-ReportFigures.py"

WD_FIND_CONTINUE = 1
WD_FIND_STOP = 0
WD_REPLACE_ALL = 2
WD_SECTION_BREAK_NEXT_PAGE = 2
WD_ALIGN_PARAGRAPH_CENTER = 1
WD_ALIGN_PARAGRAPH_JUSTIFY = 3
WD_ALIGN_PAGE_NUMBER_CENTER = 1
WD_PAGE_NUMBER_STYLE_ARABIC = 0
WD_PAGE_NUMBER_STYLE_LOWERCASE_ROMAN = 2
WD_LINE_SPACE_MULTIPLE = 5
WD_EXPORT_FORMAT_PDF = 17
WD_PAPER_A4 = 7
WD_STATISTIC_WORDS = 0
WD_STATISTIC_PAGES = 2
WD_STYLE_TYPE_PARAGRAPH = 1
WD_FIELD_EMPTY = -1


CONTENT_TYPES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Default Extension="png" ContentType="image/png"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
  <Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>
  <Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>
  <Override PartName="/word/webSettings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.webSettings+xml"/>
  <Override PartName="/word/fontTable.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml"/>
  <Override PartName="/word/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>
  <Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>
  <Override PartName="/word/endnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>
  <Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>
  <Override PartName="/word/header2.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>
  <Override PartName="/word/header3.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>
  <Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>
  <Override PartName="/word/footer2.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>
  <Override PartName="/word/footer3.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>
"""


def ensure_content_types(path: Path) -> None:
    with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_DEFLATED) as package:
        if "[Content_Types].xml" not in package.namelist():
            package.writestr("[Content_Types].xml", CONTENT_TYPES_XML)


def rebuild_as_valid_docx(path: Path) -> None:
    """Rebuild the inherited package using standards-compliant python-docx parts.

    The original generated file omitted a mandatory OPC manifest and also contains
    XML that desktop Word refuses to repair. Recreating its paragraphs and tables
    in a fresh package preserves the report content while producing a dependable
    Word input for the final institutional formatting pass.
    """

    source = Document(path)
    rebuilt = Document()

    body = rebuilt._element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)

    style_map = {
        "Titre": "Title",
        "Sous-titre": "Subtitle",
        "Titre1": "Heading 1",
        "Titre2": "Heading 2",
        "Titre3": "Heading 3",
        "Lgende": "Caption",
        "Paragraphedeliste": "List Bullet",
        "Corpsdetexte": "Normal",
        "NormalWeb": "Normal",
        "Tabledesillustrations": "Normal",
    }

    for child in source.element.body.iterchildren():
        if isinstance(child, CT_P):
            source_paragraph = Paragraph(child, source)
            source_style = ""
            try:
                source_style = source_paragraph.style.style_id
            except Exception:
                pass
            target_style = style_map.get(source_style, "Normal")
            target_paragraph = rebuilt.add_paragraph(style=target_style)

            has_drawing = bool(child.xpath(".//*[local-name()='drawing']"))
            if has_drawing:
                blips = child.xpath(".//*[local-name()='blip']")
                if blips:
                    relationship_id = blips[0].get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                    )
                    if relationship_id and relationship_id in source.part.related_parts:
                        image_blob = source.part.related_parts[relationship_id].blob
                        target_paragraph.add_run().add_picture(
                            io.BytesIO(image_blob), width=Inches(6.2)
                        )
                        target_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if source_paragraph.text:
                target_paragraph.add_run(source_paragraph.text)

            if "w:type=\"page\"" in child.xml:
                target_paragraph.add_run().add_break(7)

        elif isinstance(child, CT_Tbl):
            source_table = Table(child, source)
            row_count = len(source_table.rows)
            column_count = max((len(row.cells) for row in source_table.rows), default=1)
            target_table = rebuilt.add_table(rows=row_count, cols=column_count)
            target_table.style = "Table Grid"
            for row_index, source_row in enumerate(source_table.rows):
                for column_index, source_cell in enumerate(source_row.cells):
                    target_table.cell(row_index, column_index).text = source_cell.text

    rebuilt.core_properties.title = "Rapport de stage — Display Control"
    rebuilt.core_properties.author = "Nizar NEGAGZA"
    rebuilt.save(path)


def clean_paragraph_text(paragraph) -> str:
    return paragraph.Range.Text.rstrip("\r\x07")


def find_paragraph(document, text: str, *, last: bool = False):
    needle = text if len(text) <= 250 else text[:200]
    search_range = document.Content.Duplicate
    finder = search_range.Find
    finder.ClearFormatting()
    found = finder.Execute(
        FindText=needle,
        Forward=True,
        Wrap=WD_FIND_STOP,
        Format=False,
    )
    if found:
        paragraph = search_range.Paragraphs.Item(1)
        if clean_paragraph_text(paragraph).strip() == text:
            return paragraph
    indexes = range(document.Paragraphs.Count, 0, -1) if last else range(1, document.Paragraphs.Count + 1)
    for index in indexes:
        paragraph = document.Paragraphs.Item(index)
        if clean_paragraph_text(paragraph).strip() == text:
            return paragraph
    raise RuntimeError(f"Paragraphe introuvable : {text}")


def find_caption_paragraph(document, text: str):
    search_range = document.Content.Duplicate
    finder = search_range.Find
    finder.ClearFormatting()
    finder.Style = document.Styles("Légende")
    found = finder.Execute(
        FindText=text,
        Forward=False,
        Wrap=WD_FIND_STOP,
        Format=True,
    )
    if found:
        return search_range.Paragraphs.Item(1)
    return find_paragraph(document, text, last=True)


def find_heading_paragraph(document, text: str, *, last: bool = False):
    search_range = document.Content.Duplicate
    finder = search_range.Find
    finder.ClearFormatting()
    finder.Style = document.Styles("Titre 1")
    found = finder.Execute(
        FindText=text,
        Forward=not last,
        Wrap=WD_FIND_STOP,
        Format=True,
    )
    if found:
        return search_range.Paragraphs.Item(1)
    raise RuntimeError(f"Titre de niveau 1 introuvable : {text}")


def paragraph_exists(document, text: str) -> bool:
    try:
        find_paragraph(document, text)
        return True
    except RuntimeError:
        return False


def replace_all(document, old: str, new: str) -> None:
    search_range = document.Content
    finder = search_range.Find
    finder.ClearFormatting()
    finder.Replacement.ClearFormatting()
    changed = finder.Execute(
        FindText=old,
        ReplaceWith=new,
        Forward=True,
        Wrap=WD_FIND_CONTINUE,
        Format=False,
        Replace=WD_REPLACE_ALL,
    )
    if not changed:
        raise RuntimeError(f"Texte introuvable pour remplacement : {old}")


def replace_paragraph(document, old: str, new: str) -> None:
    paragraph = find_paragraph(document, old)
    paragraph.Range.Text = new + "\r"


def insert_block_before(document, before_title: str, paragraphs: list[str]) -> None:
    destination = find_paragraph(document, before_title, last=True)
    start = destination.Range.Start
    inserted_text = "\r".join(paragraphs) + "\r"
    insertion = document.Range(start, start)
    insertion.InsertBefore(inserted_text)
    document.Range(start, start + len(inserted_text)).Style = document.Styles("Normal")


def append_block(document, paragraphs: list[str]) -> None:
    start = document.Content.End - 1
    inserted_text = "\r" + "\r".join(paragraphs) + "\r"
    insertion = document.Range(start, start)
    insertion.InsertAfter(inserted_text)
    document.Range(start, start + len(inserted_text)).Style = document.Styles("Normal")


def insert_picture_before_caption(word, document, caption_text: str, image_path: Path) -> None:
    caption = find_caption_paragraph(document, caption_text)
    caption.Style = document.Styles("Légende")
    picture_paragraph = document.Paragraphs.Add(document.Range(caption.Range.Start, caption.Range.Start))
    picture_paragraph.Range.ParagraphFormat.Alignment = WD_ALIGN_PARAGRAPH_CENTER
    picture = picture_paragraph.Range.InlineShapes.AddPicture(
        FileName=str(image_path),
        LinkToFile=False,
        SaveWithDocument=True,
    )
    max_width = 16.2 * 28.3464567
    if picture.Width > max_width:
        picture.LockAspectRatio = True
        picture.Width = max_width


def ensure_caption_style(document, name: str) -> None:
    try:
        style = document.Styles(name)
    except Exception:
        style = document.Styles.Add(name, WD_STYLE_TYPE_PARAGRAPH)
    style.BaseStyle = document.Styles("Légende")
    style.Font.Name = "Times New Roman"
    style.Font.Size = 10
    style.Font.Italic = True
    style.ParagraphFormat.Alignment = WD_ALIGN_PARAGRAPH_CENTER
    style.ParagraphFormat.KeepWithNext = True


def convert_numbered_caption(
    document,
    old_text: str,
    *,
    kind: str,
    title: str,
    chapter: int | None = None,
) -> None:
    paragraph = find_caption_paragraph(document, old_text)
    start = paragraph.Range.Start
    if chapter is None:
        prefix = f"{kind} "
        sequence = kind
        style_name = "Légende tableau"
    else:
        prefix = f"{kind} {chapter}."
        sequence = f"FigureChapitre{chapter}"
        style_name = "Légende figure"
    paragraph.Range.Text = prefix + " — " + title + "\r"
    field_range = document.Range(start + len(prefix), start + len(prefix))
    document.Fields.Add(
        Range=field_range,
        Type=WD_FIELD_EMPTY,
        Text=f"SEQ {sequence} \\* ARABIC",
        PreserveFormatting=True,
    )
    document.Range(start, start + len(prefix) + len(title) + 8).Paragraphs.Item(1).Style = document.Styles(style_name)


def rebuild_styled_list(document, heading_text: str, next_heading_text: str, style_name: str) -> None:
    heading = find_paragraph(document, heading_text)
    next_heading = find_paragraph(document, next_heading_text)
    heading.Style = document.Styles("Titre 1")
    next_heading.Style = document.Styles("Titre 1")
    delete_end = max(heading.Range.End, next_heading.Range.Start - 1)
    document.Range(heading.Range.End, delete_end).Delete()
    find_paragraph(document, heading_text).Style = document.Styles("Titre 1")
    find_paragraph(document, next_heading_text).Style = document.Styles("Titre 1")
    heading = find_heading_paragraph(document, heading_text)
    insertion = document.Range(heading.Range.End, heading.Range.End)
    document.TablesOfContents.Add(
        Range=insertion,
        UseHeadingStyles=False,
        UpperHeadingLevel=1,
        LowerHeadingLevel=1,
        IncludePageNumbers=True,
        RightAlignPageNumbers=True,
        AddedStyles=f"{style_name},1",
        UseHyperlinks=True,
        HidePageNumbersInWeb=True,
        UseOutlineLevels=False,
    )


def ensure_abbreviations_section(document) -> None:
    try:
        find_heading_paragraph(document, "Liste des abréviations")
        return
    except RuntimeError:
        pass
    entries = [
        "API — Application Programming Interface (interface de programmation)",
        "ARM64 — Architecture ARM 64 bits",
        "CSRF — Cross-Site Request Forgery",
        "CSR — Certificate Signing Request (demande de certificat)",
        "ECDSA — Elliptic Curve Digital Signature Algorithm",
        "EF Core — Entity Framework Core",
        "HTTPS — Hypertext Transfer Protocol Secure",
        "MFA — Multi-Factor Authentication (authentification multifacteur)",
        "mTLS — Mutual Transport Layer Security",
        "RLS — Row-Level Security (sécurité au niveau des lignes)",
        "SHA-256 — Secure Hash Algorithm sur 256 bits",
        "SMTP — Simple Mail Transfer Protocol",
        "TOTP — Time-based One-Time Password",
        "UTC — Coordinated Universal Time",
    ]
    introduction = find_heading_paragraph(document, "Introduction générale")
    start = introduction.Range.Start
    text = "Liste des abréviations\r" + "\r".join(entries) + "\r"
    document.Range(start, start).InsertBefore(text)
    heading = find_paragraph(document, "Liste des abréviations", last=True)
    heading.Style = document.Styles("Titre 1")
    for entry in entries:
        find_paragraph(document, entry, last=True).Style = document.Styles("Normal")


def ensure_preliminary_list_headings(document) -> None:
    try:
        tables_heading = find_paragraph(document, "Liste des tableaux")
    except RuntimeError:
        anchor = find_paragraph(document, "Liste des abréviations", last=True)
        document.Range(anchor.Range.Start, anchor.Range.Start).InsertBefore("Liste des tableaux\r")
        tables_heading = find_paragraph(document, "Liste des tableaux")
    tables_heading.Style = document.Styles("Titre 1")

    try:
        figures_heading = find_paragraph(document, "Liste des figures")
    except RuntimeError:
        tables_heading = find_paragraph(document, "Liste des tableaux")
        document.Range(tables_heading.Range.Start, tables_heading.Range.Start).InsertBefore("Liste des figures\r")
        figures_heading = find_paragraph(document, "Liste des figures")
    figures_heading.Style = document.Styles("Titre 1")


def set_heading(document, old: str, new: str, style_name: str) -> None:
    paragraph = find_paragraph(document, old, last=True)
    paragraph.Range.Text = new + "\r"
    updated = find_paragraph(document, new, last=True)
    updated.Style = document.Styles(style_name)


def move_block(document, start_title: str, end_title: str, before_title: str) -> None:
    start_paragraph = find_paragraph(document, start_title)
    end_paragraph = find_paragraph(document, end_title)
    block = document.Range(start_paragraph.Range.Start, end_paragraph.Range.Start)
    block.Cut()
    destination = find_paragraph(document, before_title)
    insertion = document.Range(destination.Range.Start, destination.Range.Start)
    insertion.Paste()


def move_heading_section_before(document, start_title: str, before_title: str) -> None:
    start_paragraph = find_heading_paragraph(document, start_title)
    end_position = document.Content.End - 1
    found_start = False
    for index in range(1, document.Paragraphs.Count + 1):
        paragraph = document.Paragraphs.Item(index)
        if paragraph.Range.Start == start_paragraph.Range.Start:
            found_start = True
            continue
        if not found_start:
            continue
        try:
            style_name = str(paragraph.Style.NameLocal)
        except Exception:
            style_name = ""
        if style_name == "Titre 1":
            end_position = paragraph.Range.Start
            break
    block = document.Range(start_paragraph.Range.Start, end_position)
    block.Cut()
    destination = find_heading_paragraph(document, before_title)
    document.Range(destination.Range.Start, destination.Range.Start).Paste()


def remove_manual_page_break_before(document, position: int) -> None:
    start = max(0, position - 4)
    preceding = document.Range(start, position)
    text = preceding.Text
    page_break_index = text.rfind("\x0c")
    if page_break_index >= 0:
        document.Range(start + page_break_index, start + page_break_index + 1).Delete()


def insert_section_break_before(document, title: str) -> None:
    paragraph = find_heading_paragraph(document, title)
    remove_manual_page_break_before(document, paragraph.Range.Start)
    paragraph = find_heading_paragraph(document, title)
    insertion = document.Range(paragraph.Range.Start, paragraph.Range.Start)
    insertion.InsertBreak(WD_SECTION_BREAK_NEXT_PAGE)


def ensure_three_report_sections(document) -> None:
    dedication = find_heading_paragraph(document, "Dédicace")
    introduction = find_heading_paragraph(document, "Introduction générale")
    while document.Sections.Count < 3:
        dedication_section = dedication.Range.Sections.Item(1).Index
        introduction_section = introduction.Range.Sections.Item(1).Index
        if dedication_section == 1:
            position = dedication.Range.Start
        elif introduction_section == dedication_section:
            position = introduction.Range.Start
        else:
            position = introduction.Range.Start
        document.Range(position, position).InsertBreak(WD_SECTION_BREAK_NEXT_PAGE)
        dedication = find_heading_paragraph(document, "Dédicace")
        introduction = find_heading_paragraph(document, "Introduction générale")
    if document.Sections.Count != 3:
        raise RuntimeError(f"Trois sections attendues, {document.Sections.Count} trouvées.")


def normalize_final_heading_structure(document) -> None:
    final_changes = [
        ("Chapitre 3 : Réalisation de la solution", "Chapitre 3 : Réalisation, tests et résultats", "Titre 1"),
        ("3.12 Conclusion", "3.12 Synthèse de la réalisation", "Titre 2"),
        ("Chapitre 4 : Tests, résultats et bilan", "3.13 Stratégie de tests et résultats", "Titre 2"),
        ("4.1 Introduction", "3.13.1 Introduction", "Titre 3"),
        ("4.2 Stratégie de vérification", "3.13.2 Stratégie de vérification", "Titre 3"),
        ("4.3 Test du Raspberry Pi virtuel", "3.13.3 Test du Raspberry Pi virtuel", "Titre 3"),
        ("4.4 Résultats observés", "3.13.4 Résultats observés", "Titre 3"),
        ("4.5 Anomalies et limites découvertes", "3.13.5 Anomalies corrigées et limites ouvertes", "Titre 3"),
        ("4.6 Validation matérielle restant à réaliser", "3.13.6 Validation matérielle restant à réaliser", "Titre 3"),
        ("4.7 Apports du stage", "3.13.7 Apports du stage", "Titre 3"),
        ("4.8 Conclusion", "3.14 Conclusion", "Titre 2"),
    ]
    for old, new, style in final_changes:
        try:
            set_heading(document, old, new, style)
        except RuntimeError:
            # Le titre peut déjà avoir été normalisé lors dʼun passage antérieur.
            paragraph = find_paragraph(document, new, last=True)
            paragraph.Style = document.Styles(style)

    inherited_heading_texts = [
        "À mes parents et à ma famille, pour leur soutien constant, leur confiance et leurs encouragements tout au long de mon parcours.",
        "À mes enseignants, à mes encadrants et à toutes les personnes qui ont contribué à ma formation et à la réussite de ce projet.",
        "La vérification locale du 30 août 2026 totalise 102 tests .NET réussis — 44 tests de domaine, 12 tests de lʼagent et 46 tests dʼintégration —, 8 tests React et 5 scénarios Playwright/axe. La compilation Release se termine sans avertissement ni erreur, et les audits NuGet et npm ne signalent aucune vulnérabilité connue. Ces preuves valident le logiciel dans lʼenvironnement local ; elles ne remplacent ni les essais sur Raspberry Pi physique ni la validation dʼune infrastructure de recette.",
        "[8] NEGAGZA Nizar, « Documentation interne de Display Control : exigences, architecture, modèle de données, API, modèle de menace, stratégie de test et procédures de déploiement », Elastomer Solutions Maroc, Tanger, 2026.",
    ]
    for text in inherited_heading_texts:
        try:
            find_paragraph(document, text, last=True).Style = document.Styles("Normal")
        except RuntimeError:
            pass


def apply_document_formatting(word, document) -> None:
    cm = 28.3464567
    for section_index in range(1, document.Sections.Count + 1):
        section = document.Sections.Item(section_index)
        setup = section.PageSetup
        setup.PaperSize = WD_PAPER_A4
        setup.LeftMargin = 2 * cm
        setup.RightMargin = 2 * cm
        setup.TopMargin = 1.5 * cm
        setup.BottomMargin = 1.5 * cm
        setup.FooterDistance = 1 * cm
        setup.Gutter = 0
        setup.MirrorMargins = False

    def format_style(
        name: str,
        size: float,
        *,
        bold: bool = False,
        italic: bool = False,
        alignment: int = WD_ALIGN_PARAGRAPH_JUSTIFY,
        before: float = 0,
        after: float = 6,
        keep_with_next: bool = False,
        page_break_before: bool = False,
    ) -> None:
        style = document.Styles(name)
        style.Font.Name = "Times New Roman"
        style.Font.Size = size
        style.Font.Bold = bold
        style.Font.Italic = italic
        style.ParagraphFormat.Alignment = alignment
        style.ParagraphFormat.LineSpacingRule = WD_LINE_SPACE_MULTIPLE
        style.ParagraphFormat.LineSpacing = size * 1.15
        style.ParagraphFormat.SpaceBefore = before
        style.ParagraphFormat.SpaceAfter = after
        style.ParagraphFormat.KeepWithNext = keep_with_next
        style.ParagraphFormat.PageBreakBefore = page_break_before

    format_style("Normal", 12)
    format_style("Liste à puces", 12)
    format_style(
        "Titre 1",
        16,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH_CENTER,
        after=12,
        page_break_before=True,
    )
    format_style("Titre 2", 14, bold=True, before=12, keep_with_next=True)
    format_style("Titre 3", 12, bold=True, before=8, after=4, keep_with_next=True)
    format_style(
        "Titre",
        20,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH_CENTER,
        after=12,
    )
    format_style(
        "Sous-titre",
        16,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH_CENTER,
        after=10,
    )
    format_style(
        "Légende",
        10,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH_CENTER,
        after=6,
        keep_with_next=True,
    )

    for table_index in range(1, document.Tables.Count + 1):
        table = document.Tables.Item(table_index)
        table.Rows.Alignment = WD_ALIGN_PARAGRAPH_CENTER
        table.Range.Font.Name = "Times New Roman"
        table.Range.Font.Size = 10.5
        table.Range.ParagraphFormat.LineSpacingRule = WD_LINE_SPACE_MULTIPLE
        table.Range.ParagraphFormat.LineSpacing = 10.5 * 1.15

    for shape_index in range(1, document.InlineShapes.Count + 1):
        shape = document.InlineShapes.Item(shape_index)
        shape.Range.ParagraphFormat.Alignment = WD_ALIGN_PARAGRAPH_CENTER


def configure_page_numbers(document) -> None:
    if document.Sections.Count != 3:
        raise RuntimeError(
            f"Trois sections attendues (garde, préliminaires, rapport), "
            f"{document.Sections.Count} trouvées."
        )

    for section_index in range(1, 4):
        section = document.Sections.Item(section_index)
        section.PageSetup.OddAndEvenPagesHeaderFooter = False
        section.PageSetup.DifferentFirstPageHeaderFooter = False
        footer = section.Footers.Item(1)
        footer.LinkToPrevious = False
        while footer.PageNumbers.Count:
            footer.PageNumbers.Item(1).Delete()

    preliminary_footer = document.Sections.Item(2).Footers.Item(1)
    preliminary_numbers = preliminary_footer.PageNumbers
    preliminary_numbers.RestartNumberingAtSection = True
    preliminary_numbers.StartingNumber = 1
    preliminary_numbers.NumberStyle = WD_PAGE_NUMBER_STYLE_LOWERCASE_ROMAN
    preliminary_numbers.Add(WD_ALIGN_PAGE_NUMBER_CENTER, True)

    report_footer = document.Sections.Item(3).Footers.Item(1)
    report_numbers = report_footer.PageNumbers
    report_numbers.RestartNumberingAtSection = True
    report_numbers.StartingNumber = 1
    report_numbers.NumberStyle = WD_PAGE_NUMBER_STYLE_ARABIC
    report_numbers.Add(WD_ALIGN_PAGE_NUMBER_CENTER, True)


def rebuild_table_of_contents(document) -> None:
    while document.TablesOfContents.Count:
        document.TablesOfContents.Item(1).Delete()

    heading = find_heading_paragraph(document, "Table des matières")
    next_heading = find_heading_paragraph(document, "Liste des figures")
    old_body = document.Range(heading.Range.End, next_heading.Range.Start)
    old_body.Delete()
    heading = find_heading_paragraph(document, "Table des matières")
    insertion = document.Range(heading.Range.End, heading.Range.End)
    document.TablesOfContents.Add(insertion)


def clear_manual_table_of_contents_body(document) -> None:
    heading = find_heading_paragraph(document, "Table des matières")
    next_heading = find_heading_paragraph(document, "Liste des figures")
    document.Range(heading.Range.End, next_heading.Range.Start).Delete()
    find_paragraph(document, "Table des matières").Style = document.Styles("Titre 1")
    find_paragraph(document, "Liste des figures").Style = document.Styles("Titre 1")


def set_properties(document) -> None:
    properties = {
        "Title": "Rapport de stage / PFE — Display Control",
        "Author": "M. NEGAGZA Nizar",
        "Subject": "Conception et développement d’une plateforme sécurisée de gestion d’affichage sur Raspberry Pi",
        "Keywords": "Display Control, Raspberry Pi, ASP.NET Core, React, PostgreSQL, mTLS, stage, PFE",
    }
    for name, value in properties.items():
        try:
            document.BuiltInDocumentProperties(name).Value = value
        except Exception:
            pass


def build_report() -> dict[str, object]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    subprocess.run(["python", str(FIGURE_GENERATOR)], cwd=ROOT, check=True)
    shutil.copy2(SOURCE, DOCX_OUTPUT)
    ensure_content_types(DOCX_OUTPUT)
    rebuild_as_valid_docx(DOCX_OUTPUT)
    if PDF_OUTPUT.exists():
        PDF_OUTPUT.unlink()

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    word.ScreenUpdating = False
    document = None

    try:
        document = word.Documents.Open(
            FileName=str(DOCX_OUTPUT),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
            OpenAndRepair=True,
        )
        print("Word : document ouvert", flush=True)

        # Page de garde et informations administratives confirmées.
        replace_all(document, "RAPPORT DE STAGE", "RAPPORT DE STAGE / PROJET DE FIN D’ÉTUDES")
        replace_all(document, "[Nom et prénom de lʼétudiant]", "M. NEGAGZA Nizar")
        replace_all(document, "ELASTOMER SOLUTIONS — TANGER", "ELASTOMER SOLUTIONS MAROC SARL — TANGER")
        replace_all(
            document,
            "Période : août – septembre 2026  |  Durée : six semaines",
            "Période : du 3 août au 11 septembre 2026  |  Durée : six semaines",
        )
        replace_all(
            document,
            "Document préparé le 27 août 2026",
            "Soutenance : [À compléter]  |  Année universitaire : 2025–2026",
        )

        administrative_table = document.Tables.Item(1)
        administrative_table.Cell(1, 1).Range.Text = "Information"
        administrative_table.Cell(1, 2).Range.Text = "Détail"
        administrative_table.Cell(2, 1).Range.Text = "Présenté par"
        administrative_table.Cell(2, 2).Range.Text = "M. NEGAGZA Nizar"
        administrative_table.Cell(3, 1).Range.Text = "Établissement et filière"
        administrative_table.Cell(3, 2).Range.Text = "ENSI Tanger — [À compléter]"
        administrative_table.Cell(4, 1).Range.Text = "Diplôme préparé"
        administrative_table.Cell(4, 2).Range.Text = "Ingénieur d’État en Génie [À compléter]"
        administrative_table.Cell(5, 1).Range.Text = "Encadrant professionnel"
        administrative_table.Cell(5, 2).Range.Text = "[À compléter]"
        administrative_table.Cell(6, 1).Range.Text = "Encadrant pédagogique"
        administrative_table.Cell(6, 2).Range.Text = "[À compléter]"
        for label, value in [
            ("Service d’accueil", "[À compléter]"),
            ("Date de soutenance", "[À compléter]"),
            ("Membres du jury", "[À compléter]"),
            ("Année universitaire", "2025–2026"),
        ]:
            row = administrative_table.Rows.Add()
            row.Cells(1).Range.Text = label
            row.Cells(2).Range.Text = value

        # Dédicace obligatoire ENSI.
        if not paragraph_exists(document, "Dédicace"):
            thanks = find_paragraph(document, "Remerciements")
            insertion_position = thanks.Range.Start
            dedication_text = (
                "Dédicace\r"
                "À mes parents et à ma famille, pour leur soutien constant, leur confiance et leurs encouragements tout au long de mon parcours.\r"
                "À mes enseignants, à mes encadrants et à toutes les personnes qui ont contribué à ma formation et à la réussite de ce projet.\r"
                "\x0c"
            )
            document.Range(insertion_position, insertion_position).InsertBefore(dedication_text)
            dedication = find_paragraph(document, "Dédicace")
            dedication.Style = thanks.Style
            dedication.Range.ParagraphFormat.Alignment = WD_ALIGN_PARAGRAPH_CENTER

        # Ordre exigé : dédicace, remerciements, résumé, abstract, sommaire,
        # figures, tableaux, abréviations, introduction générale.
        # Réordonner les quatre rubriques préliminaires de manière stable,
        # indépendamment de leur ordre dans le document hérité.
        move_heading_section_before(document, "Liste des acronymes", "Introduction générale")
        move_heading_section_before(document, "Liste des tableaux", "Liste des acronymes")
        move_heading_section_before(document, "Liste des figures", "Liste des tableaux")
        move_heading_section_before(document, "Table des matières", "Liste des figures")
        replace_all(document, "Liste des acronymes", "Liste des abréviations")
        clear_manual_table_of_contents_body(document)
        print("Word : préliminaires réordonnés", flush=True)

        # Présentation de l’entreprise issue de son site officiel.
        replace_paragraph(
            document,
            "Le stage est réalisé au sein dʼElastomer Solutions à Tanger. Dans ce rapport, la présentation de lʼentreprise reste volontairement limitée aux informations confirmées par le stagiaire : son nom, son implantation à Tanger et son rôle dʼorganisme dʼaccueil. Les données institutionnelles détaillées, lʼorganigramme, les chiffres dʼactivité et les éléments confidentiels devront être ajoutés uniquement après validation par lʼentreprise.",
            "Le stage est réalisé au sein dʼElastomer Solutions Maroc SARL, implantée dans la Zone Franche dʼExportation de Tanger. Le groupe développe et fabrique des composants en caoutchouc, en thermoplastique et des pièces bi-matière caoutchouc-plastique destinés principalement à lʼindustrie automobile. Le site marocain a démarré lʼinjection de composants en caoutchouc en 2013 ; une nouvelle unité de production à Tanger est entrée en activité en 2016 afin dʼaccompagner la croissance du groupe et de servir les marchés dʼEurope, dʼAfrique et des Amériques [1].",
        )
        replace_all(document, "Information à compléter avant dépôt final", "Repères institutionnels vérifiés")
        replace_all(
            document,
            "Ajouter la présentation officielle dʼElastomer Solutions, son secteur dʼactivité exact, son organisation et le service dʼaccueil après validation par lʼencadrant. Aucune donnée non confirmée nʼa été inventée dans cette version.",
            "Elastomer Solutions Maroc SARL — industrie automobile — Zone Franche dʼExportation, îlot 64, lot 3, 90100 Tanger. Source : site officiel dʼElastomer Solutions, consulté le 31 août 2026. Le service dʼaccueil et lʼorganigramme restent à confirmer avec lʼencadrant professionnel.",
        )

        # Trois chapitres principaux selon la structure ENSI.
        replace_all(
            document,
            "Chapitre 3 : Réalisation de la solution",
            "Chapitre 3 : Réalisation, tests et résultats",
        )
        replace_all(document, "3.12 Conclusion", "3.12 Synthèse de la réalisation")
        set_heading(
            document,
            "Chapitre 4 : Tests, résultats et bilan",
            "3.13 Stratégie de tests et résultats",
            "Titre 2",
        )
        heading_changes = {
            "4.1 Introduction": "3.13.1 Introduction",
            "4.2 Stratégie de vérification": "3.13.2 Stratégie de vérification",
            "4.3 Test du Raspberry Pi virtuel": "3.13.3 Test du Raspberry Pi virtuel",
            "4.4 Résultats observés": "3.13.4 Résultats observés",
            "4.5 Anomalies et limites découvertes": "3.13.5 Anomalies corrigées et limites ouvertes",
            "4.6 Validation matérielle restant à réaliser": "3.13.6 Validation matérielle restant à réaliser",
            "4.7 Apports du stage": "3.13.7 Apports du stage",
        }
        for old, new in heading_changes.items():
            set_heading(document, old, new, "Titre 3")
        set_heading(document, "4.8 Conclusion", "3.14 Conclusion", "Titre 2")
        print("Word : structure à trois chapitres appliquée", flush=True)

        # Mise à jour factuelle au dernier contrôle disponible.
        replace_all(
            document,
            "Tableau 13 — Résultats du test virtuel du 27 août 2026",
            "Tableau 13 — Résultats de la campagne locale du 30 août 2026",
        )
        replace_paragraph(
            document,
            "Le test a identifié un défaut du script Windows Start-VirtualDevice.ps1 : lorsquʼil utilise dotnet run depuis le dossier source, lʼAPI locale du lecteur fonctionne mais les fichiers React compilés sont recherchés dans un autre répertoire, ce qui provoque une réponse 404 sur la racine. Lʼexécution depuis le dossier compilé, équivalente au paquet publié sur le Pi, sert correctement lʼinterface. Le script de simulation doit être ajusté pour reproduire ce comportement sans manipulation manuelle.",
            "Le premier essai virtuel a révélé un défaut de chemin dans le lanceur Windows : lʼAPI locale du lecteur fonctionnait, mais les fichiers React compilés étaient recherchés dans un autre répertoire. Le lanceur et la préparation de lʼenvironnement ont ensuite été durcis ; le parcours complet est désormais reproductible depuis le paquet compilé, qui correspond au mode dʼexécution prévu sur le Raspberry Pi.",
        )
        replace_paragraph(
            document,
            "Un scénario dʼintégration complet a également rencontré une attente devenue obsolète après le renforcement de la politique de mot de passe. Le mot de passe de test contenait le mot « viewer », identique à la partie locale de lʼadresse électronique, et il a été correctement refusé comme information personnelle. Le comportement de sécurité est cohérent, mais la donnée du test doit être mise à jour avant de compter ce scénario comme vert dans la vérification actuelle.",
            "Un scénario dʼintégration a également mis en évidence une donnée de test devenue incompatible avec la politique de mot de passe renforcée. La fixture a été corrigée sans affaiblir la règle de sécurité, puis toute la campagne a été relancée. Ce cas confirme lʼintérêt de conserver des tests représentatifs des politiques réellement appliquées.",
        )
        replace_paragraph(
            document,
            "Ces résultats illustrent lʼintérêt dʼun rapport honnête. Les preuves antérieures du dépôt indiquent une campagne complète réussie avant les derniers changements. Après ces changements, la suite agent passe et le parcours virtuel fonctionne, mais la campagne complète doit être relancée après correction de la fixture de test. Aucune affirmation de préparation à la production ne doit reposer uniquement sur le test virtuel.",
            "La vérification locale du 30 août 2026 totalise 102 tests .NET réussis — 44 tests de domaine, 12 tests de lʼagent et 46 tests dʼintégration —, 8 tests React et 5 scénarios Playwright/axe. La compilation Release se termine sans avertissement ni erreur, et les audits NuGet et npm ne signalent aucune vulnérabilité connue. Ces preuves valident le logiciel dans lʼenvironnement local ; elles ne remplacent ni les essais sur Raspberry Pi physique ni la validation dʼune infrastructure de recette.",
        )
        replace_paragraph(
            document,
            "Le stage a abouti à un socle fonctionnel important : administration multi-client, authentification avec MFA, enrôlement des appareils, inventaire, cycle de licence, contenus versionnés, playlists, synchronisation et lecteur local. Le test virtuel a confirmé le chemin critique et le comportement pendant une coupure. La solution est donc démontrable et prête pour une campagne matérielle, mais elle nʼest pas encore déclarée prête pour la production.",
            "Le stage a abouti à un socle fonctionnel important : administration multi-client, authentification avec MFA, enrôlement des appareils, inventaire, cycle de licence, contenus versionnés, playlists, synchronisation et lecteur local. Le parcours virtuel ainsi que 102 tests .NET, 8 tests React et 5 scénarios navigateur ont confirmé le chemin critique et plusieurs comportements de panne. La solution est démontrable et prête pour une campagne matérielle, mais elle nʼest pas encore déclarée prête pour la production.",
        )
        replace_paragraph(
            document,
            "À court terme, les priorités sont la correction du lanceur virtuel, lʼactualisation de la fixture de mot de passe et lʼexécution sur un Raspberry Pi connecté à un écran. À moyen terme, il faudra préparer un environnement de recette avec DNS, TLS public, stockage privé, SMTP, supervision et sauvegardes, puis conduire les tests de charge, dʼaccessibilité manuelle et de sécurité. Enfin, un mécanisme complet de mise à jour signée et de retour arrière renforcera lʼexploitation dʼune flotte dʼappareils.",
            "À court terme, la priorité est lʼexécution du paquet ARM64 sur un Raspberry Pi 4 ou 5 connecté à un écran, avec validation du numéro de série, de lʼaffichage HDMI, de Chromium kiosk, de systemd et des scénarios de coupure. À moyen terme, il faudra préparer un environnement de recette avec DNS, TLS public, stockage privé, SMTP, supervision et sauvegardes, puis conduire les tests de charge, dʼaccessibilité manuelle et de sécurité. Enfin, un mécanisme complet de mise à jour signée et de retour arrière renforcera lʼexploitation dʼune flotte dʼappareils.",
        )
        replace_all(document, "État au 27 août 2026", "État au 31 août 2026")
        replace_all(document, "Tests agent actuels", "Campagne logicielle du 30 août 2026")
        replace_all(document, "12/12 réussis", "102/102 .NET ; 8/8 React ; 5/5 navigateur")
        replace_all(
            document,
            "Suite exécutée après le test virtuel",
            "Compilation, intégration, interfaces et accessibilité automatisée",
        )
        replace_all(document, "Générés par ce document", "Générés le 31 août 2026")

        # Références numérotées et webographie datée.
        replace_paragraph(
            document,
            "1. Microsoft, documentation ASP.NET Core et Entity Framework Core, versions .NET 10.",
            "[1] Elastomer Solutions, « About us — Morocco production plant », https://www.elastomer-solutions.com/about-us/ (consulté le 31/08/2026).\r[2] Microsoft, « Documentation ASP.NET Core et Entity Framework Core », https://learn.microsoft.com/aspnet/core/ et https://learn.microsoft.com/ef/core/ (consultés le 31/08/2026).",
        )
        replace_paragraph(
            document,
            "2. PostgreSQL Global Development Group, documentation PostgreSQL 18 — Row Security Policies.",
            "[3] PostgreSQL Global Development Group, « Row Security Policies — PostgreSQL 18 », https://www.postgresql.org/docs/18/ddl-rowsecurity.html (consulté le 31/08/2026).",
        )
        replace_paragraph(
            document,
            "3. React Documentation, React 19 ; Vite Documentation, Vite 8.",
            "[4] React, « React Documentation », https://react.dev/ ; Vite, « Vite Guide », https://vite.dev/guide/ (consultés le 31/08/2026).",
        )
        replace_paragraph(
            document,
            "4. OWASP, Application Security Verification Standard 5.0 et recommandations CSRF/authentification.",
            "[5] OWASP Foundation, « Application Security Verification Standard 5.0 », https://owasp.org/www-project-application-security-verification-standard/ (consulté le 31/08/2026).",
        )
        replace_paragraph(
            document,
            "5. IETF, TLS 1.3, certificats X.509 et JSON Web Signature ES256.",
            "[6] IETF, « RFC 8446 — The Transport Layer Security (TLS) Protocol Version 1.3 », 2018, https://www.rfc-editor.org/rfc/rfc8446 (consulté le 31/08/2026).",
        )
        replace_paragraph(
            document,
            "6. Raspberry Pi Documentation, Raspberry Pi OS, Chromium et configuration de lʼaffichage.",
            "[7] Raspberry Pi Ltd., « Raspberry Pi Documentation », https://www.raspberrypi.com/documentation/ (consulté le 31/08/2026).",
        )
        replace_paragraph(
            document,
            "7. Documentation interne du projet : exigences, architecture, modèle de données, API, modèle de menace, stratégie de test et procédures de déploiement.",
            "[8] NEGAGZA Nizar, « Documentation interne de Display Control : exigences, architecture, modèle de données, API, modèle de menace, stratégie de test et procédures de déploiement », Elastomer Solutions Maroc, Tanger, 2026.",
        )

        # Citations de cadrage : chaque référence de la bibliographie est ainsi
        # reliée à un passage explicite du corps du rapport.
        insert_block_before(
            document,
            "3.3 Organisation du code",
            [
                "Les choix de mise en œuvre ont été confrontés aux documentations officielles dʼASP.NET Core et dʼEntity Framework Core [2], aux politiques de sécurité des lignes de PostgreSQL 18 [3], ainsi quʼaux guides React et Vite [4]. Les exigences de sécurité sʼappuient sur lʼOWASP ASVS [5] et sur les propriétés de TLS 1.3 décrites par la RFC 8446 [6]. La préparation du matériel reprend les recommandations publiques de Raspberry Pi Ltd. [7], tandis que les décisions propres au projet sont consignées dans la documentation interne versionnée [8].",
            ],
        )

        # Figures exigées par la trame ENSI. Les textes dʼintroduction et
        # dʼinterprétation évitent les schémas décoratifs non commentés.
        figure_insertions = [
            (
                "1.6 Résultats attendus et critères de réussite",
                "La figure temporaire 1.1 traduit la planification hebdomadaire en une vue synthétique. Elle montre que la sécurité, les tests et la rédaction ne sont pas relégués à la fin du stage mais progressent en parallèle de lʼimplémentation.",
                "Figure temporaire 1.1 — Planification du stage ou diagramme de Gantt",
                "La superposition des activités reflète une démarche incrémentale : chaque tranche fonctionnelle est accompagnée de sa documentation et de ses preuves locales. Source : Élaboration personnelle.",
            ),
            (
                "2.4 Besoins fonctionnels",
                "La figure temporaire 2.1 met en relation les acteurs avec leurs principaux cas dʼutilisation. Elle souligne la séparation entre administration de plateforme, administration dʼun tenant, gestion éditoriale et protocole autonome de lʼappareil.",
                "Figure temporaire 2.1 — Diagramme global des cas dʼutilisation",
                "Cette séparation limite les privilèges : un gestionnaire de contenu nʼobtient pas automatiquement le pouvoir de transférer une licence, tandis que le Raspberry Pi nʼutilise jamais une session humaine. Source : Élaboration personnelle.",
            ),
            (
                "2.7 Modèle de données",
                "La figure temporaire 2.3 détaille lʼorganisation interne du monolithe modulaire présenté précédemment. Les dépendances convergent vers les contrats applicatifs et les règles du domaine, alors que lʼinfrastructure fournit les adaptateurs techniques.",
                "Figure temporaire 2.3 — Architecture du monolithe modulaire",
                "Cette organisation conserve une unité de déploiement unique tout en permettant de tester les invariants sans démarrer lʼAPI ni PostgreSQL. Elle rend aussi une extraction future de module possible si un besoin dʼexploitation mesuré lʼexige. Source : Élaboration personnelle.",
            ),
            (
                "2.8 Identité du Raspberry Pi",
                "La figure temporaire 2.4 synthétise les entités structurantes et leurs relations. Elle ne remplace pas le modèle EF Core complet, mais rend visibles les liens qui doivent conserver le même tenant.",
                "Figure temporaire 2.4 — Modèle conceptuel ou diagramme de données simplifié",
                "Les clés étrangères composites incluant tenant_id complètent la RLS : elles empêchent quʼune erreur applicative relie une licence, une affectation ou un contenu à un appareil dʼun autre client. Source : Élaboration personnelle.",
            ),
            (
                "2.11 Conclusion",
                "La figure temporaire 2.7 représente les frontières de confiance traversées par les utilisateurs et les appareils. Une preuve différente est exigée à chaque frontière : session et MFA pour lʼhumain, mTLS pour lʼappareil, RLS pour les données et condensat pour le contenu.",
                "Figure temporaire 2.7 — Frontières de confiance et mécanismes de sécurité",
                "Le comportement de refus par défaut relie ces mécanismes : si lʼune des preuves devient absente, expirée ou incohérente, le lecteur adopte un état sûr. Cette politique ne prétend toutefois pas résister absolument à un attaquant disposant dʼun accès physique et root. Source : Élaboration personnelle.",
            ),
            (
                "3.4 Backend ASP.NET Core",
                "La figure temporaire 3.1 présente les principaux répertoires du monorepo. Cette organisation rapproche les applications et leurs tests tout en séparant clairement code métier, déploiement, scripts et documentation.",
                "Figure temporaire 3.1 — Organisation du dépôt source",
                "Le découpage facilite la traçabilité : les exigences et décisions sont conservées dans docs, les automatismes reproductibles dans scripts, et les preuves exécutables dans tests. Source : Élaboration personnelle.",
            ),
            (
                "3.9 Contenus, playlists et synchronisation",
                "La figure temporaire 3.2 suit un heartbeat depuis lʼagent jusquʼà la décision renvoyée. Le certificat TLS est validé avant le workflow métier, puis lʼopération est exécutée dans une transaction portant le contexte du tenant.",
                "Figure temporaire 3.2 — Flux dʼun heartbeat mTLS",
                "La réponse contient un état explicite et, lorsque toutes les conditions sont satisfaites, les informations nécessaires à la licence et à la synchronisation. Lʼadresse IP ou MAC observée reste un attribut dʼinventaire, jamais la preuve principale dʼidentité. Source : Élaboration personnelle.",
            ),
            (
                "3.10 Lecteur local et états sûrs",
                "La figure temporaire 3.3 résume la chaîne de validation dʼun contenu. Plusieurs contrôles indépendants sont nécessaires avant quʼune version ne puisse être référencée par une playlist publiée.",
                "Figure temporaire 3.3 — Chaîne de validation dʼun contenu",
                "Une indisponibilité de ClamAV maintient le fichier en quarantaine. Après un résultat propre, une approbation humaine reste obligatoire ; la version publiée devient ensuite immuable et son SHA-256 accompagne le manifeste. Source : Élaboration personnelle.",
            ),
        ]
        for before_title, introduction, caption, interpretation in figure_insertions:
            insert_block_before(document, before_title, [introduction, "", caption, interpretation])
        print("Word : textes des figures insérés", flush=True)

        insert_block_before(
            document,
            "3.11 Déploiement Raspberry Pi",
            [
                "[Insérer ici une capture réelle et expurgée de lʼinterface dʼadministration]",
                "[Insérer ici une capture réelle et expurgée du lecteur local]",
                "Ces emplacements sont volontairement laissés sans image : aucune capture fournie ne peut être inventée ni utilisée comme substitut aux preuves techniques dʼautorisation, dʼisolation ou dʼintégrité.",
            ],
        )

        # Annexes complémentaires, liste des informations manquantes et
        # checklist de conformité demandées dans le livrable.
        append_block(
            document,
            [
                "Annexe C : Architecture détaillée et dictionnaire synthétique",
                "La solution serveur est composée des projets Domain, Application, Infrastructure et Api. Domain porte les entités et invariants ; Application définit les contrats et cas dʼusage ; Infrastructure implémente EF Core, PostgreSQL, le stockage privé et les primitives cryptographiques ; Api compose les dépendances et expose les contrôleurs. DeviceAgent constitue un processus distinct installé sur le Raspberry Pi.",
                "Entités structurantes : Tenant isole une organisation ; TenantMembership relie un utilisateur à un rôle ; Device représente un écran ; DeviceCertificate porte lʼétat de lʼidentité mTLS ; DeviceLicense définit lʼautorisation temporelle ; ContentAsset et ContentVersion séparent lʼobjet éditorial de sa version immuable ; PlaylistVersion fige une publication ; DesiredState décrit lʼétat compilé destiné à un appareil.",
                "Annexe D : Synthèse des résultats de tests",
                "La campagne locale datée du 30 août 2026 a réussi 102 tests .NET, dont 44 tests de domaine, 12 tests de lʼagent et 46 tests dʼintégration. Elle a aussi réussi 8 tests React et 5 scénarios Playwright/axe. La compilation Release nʼa produit ni avertissement ni erreur et les contrôles de dépendances nʼont signalé aucune vulnérabilité connue. Ces valeurs proviennent du dossier de preuve versionné et décrivent exclusivement lʼenvironnement local.",
                "Restent à exécuter : essais Raspberry Pi 4/5 et HDMI, validation systemd et Chromium kiosk sur ARM64, endurance de 24 à 72 heures, environnement DNS/TLS public, livraison SMTP réelle, charge représentative, restauration isolée et revue de sécurité indépendante.",
                "Annexe E : Procédure résumée dʼinstallation sur Raspberry Pi",
                "1. Préparer Raspberry Pi OS 64 bits, mettre le système à jour et vérifier lʼarchitecture ARM64.",
                "2. Copier lʼarchive de déploiement et son fichier SHA-256 par un canal administré, puis vérifier lʼempreinte avant toute installation.",
                "3. Exécuter le script deploy/pi/install.sh avec les droits nécessaires ; il crée les comptes de service, les répertoires privés et une version immuable sous /opt.",
                "4. Configurer uniquement les paramètres publics requis. Le code dʼenrôlement est saisi temporairement et ne doit pas être conservé dans le rapport, les journaux ou lʼarchive.",
                "5. Activer display-control-agent.service et display-control-kiosk.service, puis contrôler leur état avec systemctl et journalctl sans exposer les secrets.",
                "6. Exécuter la fiche de recette matérielle : inventaire, enrôlement, mTLS, licence, contenu, redémarrage, coupure réseau, expiration et retour à lʼétat sûr.",
                "Annexe F : Informations à fournir et checklist ENSI",
                "Informations administratives restant à fournir : filière exacte ; intitulé complet du diplôme ; service dʼaccueil ; nom de lʼencadrant professionnel ; nom de lʼencadrant pédagogique ; date de soutenance ; membres du jury ; logos officiels validés ; éventuelles captures réelles et expurgées.",
                "Checklist — conforme : ordre des parties ; trois chapitres principaux ; format A4 ; marges 2 cm à gauche et à droite, 1,5 cm en haut et en bas ; Times New Roman 12 ; interligne 1,15 ; texte justifié ; titres hiérarchisés ; page de garde sans numéro ; préliminaires en chiffres romains ; corps en chiffres arabes recommençant à 1 ; table des matières et listes automatiques ; figures et tableaux introduits et commentés ; références numérotées ; annexes sans secret.",
                "Checklist — à finaliser avant dépôt : compléter les champs administratifs ; insérer les logos approuvés ; remplacer les deux emplacements de captures si des images réelles sont disponibles ; mettre à jour les résultats après la recette physique ; relire les coupures de pages dans la version finale de Word ; vérifier les renvois et accepter uniquement la dernière table des matières actualisée.",
            ],
        )
        for annex_heading in [
            "Annexe C : Architecture détaillée et dictionnaire synthétique",
            "Annexe D : Synthèse des résultats de tests",
            "Annexe E : Procédure résumée dʼinstallation sur Raspberry Pi",
            "Annexe F : Informations à fournir et checklist ENSI",
        ]:
            find_paragraph(document, annex_heading, last=True).Style = document.Styles("Titre 1")
        print("Word : annexes A à F complétées", flush=True)

        # Insertion des treize schémas, puis conversion des légendes en champs
        # SEQ afin que leur numérotation soit mise à jour par Word.
        figure_files = {
            "Figure temporaire 1.1 — Planification du stage ou diagramme de Gantt": "figure-1-1-gantt.png",
            "Figure temporaire 2.1 — Diagramme global des cas dʼutilisation": "figure-2-1-cas-utilisation.png",
            "Figure 1 — Architecture fonctionnelle globale de Display Control": "figure-2-2-architecture-globale.png",
            "Figure temporaire 2.3 — Architecture du monolithe modulaire": "figure-2-3-monolithe-modulaire.png",
            "Figure temporaire 2.4 — Modèle conceptuel ou diagramme de données simplifié": "figure-2-4-modele-donnees.png",
            "Figure 2 — Parcours dʼenrôlement dʼun appareil": "figure-2-5-enrolement.png",
            "Figure 3 — Décision de licence et synchronisation du contenu": "figure-2-6-licence-sync.png",
            "Figure temporaire 2.7 — Frontières de confiance et mécanismes de sécurité": "figure-2-7-frontieres-confiance.png",
            "Figure temporaire 3.1 — Organisation du dépôt source": "figure-3-1-depot.png",
            "Figure temporaire 3.2 — Flux dʼun heartbeat mTLS": "figure-3-2-heartbeat.png",
            "Figure temporaire 3.3 — Chaîne de validation dʼun contenu": "figure-3-3-validation-contenu.png",
            "Figure 4 — Frontière entre le serveur, lʼagent et le lecteur local": "figure-3-4-architecture-locale.png",
            "Figure 5 — Scénario exécuté avec le Raspberry Pi virtuel": "figure-3-5-pi-virtuel.png",
        }
        for caption_text, filename in figure_files.items():
            insert_picture_before_caption(word, document, caption_text, FIGURES_DIR / filename)
        print("Word : treize images incorporées", flush=True)

        ensure_caption_style(document, "Légende figure")
        ensure_caption_style(document, "Légende tableau")
        figure_captions = [
            ("Figure temporaire 1.1 — Planification du stage ou diagramme de Gantt", 1, "Planification du stage ou diagramme de Gantt"),
            ("Figure temporaire 2.1 — Diagramme global des cas dʼutilisation", 2, "Diagramme global des cas dʼutilisation"),
            ("Figure 1 — Architecture fonctionnelle globale de Display Control", 2, "Architecture fonctionnelle globale de Display Control"),
            ("Figure temporaire 2.3 — Architecture du monolithe modulaire", 2, "Architecture du monolithe modulaire"),
            ("Figure temporaire 2.4 — Modèle conceptuel ou diagramme de données simplifié", 2, "Modèle conceptuel ou diagramme de données simplifié"),
            ("Figure 2 — Parcours dʼenrôlement dʼun appareil", 2, "Parcours dʼenrôlement dʼun Raspberry Pi"),
            ("Figure 3 — Décision de licence et synchronisation du contenu", 2, "Décision de licence et synchronisation du contenu"),
            ("Figure temporaire 2.7 — Frontières de confiance et mécanismes de sécurité", 2, "Frontières de confiance et mécanismes de sécurité"),
            ("Figure temporaire 3.1 — Organisation du dépôt source", 3, "Organisation du dépôt source"),
            ("Figure temporaire 3.2 — Flux dʼun heartbeat mTLS", 3, "Flux dʼun heartbeat mTLS"),
            ("Figure temporaire 3.3 — Chaîne de validation dʼun contenu", 3, "Chaîne de validation dʼun contenu"),
            ("Figure 4 — Frontière entre le serveur, lʼagent et le lecteur local", 3, "Architecture locale agent–API loopback–lecteur React"),
            ("Figure 5 — Scénario exécuté avec le Raspberry Pi virtuel", 3, "Scénario du test Raspberry Pi virtuel"),
        ]
        for old_caption, chapter, title in figure_captions:
            convert_numbered_caption(document, old_caption, kind="Figure", chapter=chapter, title=title)

        table_titles = [
            "Informations administratives du stage",
            "Acronymes utilisés dans le rapport",
            "Planification prévisionnelle du stage sur six semaines",
            "Comparaison des approches de gestion dʼun écran distant",
            "Acteurs et responsabilités de Display Control",
            "Synthèse des besoins fonctionnels",
            "Besoins non fonctionnels principaux",
            "Principales entités métier",
            "Technologies et rôles dans la solution",
            "Organisation principale du dépôt",
            "Inventaire remonté par le Raspberry Pi virtuel",
            "Niveaux de test et portée des preuves",
            "Résultats de la campagne locale du 30 août 2026",
            "Synthèse de la livraison du projet",
        ]
        for number, title in enumerate(table_titles, start=1):
            old_caption = f"Tableau {number} — {title}"
            try:
                convert_numbered_caption(
                    document,
                    old_caption,
                    kind="Tableau",
                    title=title,
                )
            except RuntimeError:
                if number != 13:
                    raise
                convert_numbered_caption(
                    document,
                    "Tableau 13 — Résultats du test virtuel du 27 août 2026",
                    kind="Tableau",
                    title=title,
                )
        print("Word : légendes automatiques configurées", flush=True)

        # Sections et pagination : garde sans numéro, préliminaires en romain,
        # corps du rapport en chiffres arabes à partir de l’introduction.
        insert_section_break_before(document, "Dédicace")
        insert_section_break_before(document, "Introduction générale")
        ensure_three_report_sections(document)

        apply_document_formatting(word, document)
        rebuild_table_of_contents(document)
        ensure_abbreviations_section(document)
        ensure_preliminary_list_headings(document)
        normalize_final_heading_structure(document)
        ensure_three_report_sections(document)
        configure_page_numbers(document)
        set_properties(document)
        print("Word : sections, listes et pagination configurées", flush=True)

        document.Fields.Update()
        if document.TablesOfContents.Count:
            document.TablesOfContents.Item(1).Update()
        word.ScreenUpdating = True
        document.Repaginate()
        document.Save()
        document.ExportAsFixedFormat(str(PDF_OUTPUT), WD_EXPORT_FORMAT_PDF)
        print("Word : DOCX enregistré et PDF exporté", flush=True)
        document.Close(SaveChanges=False)
        document = None

        check = word.Documents.Open(
            FileName=str(DOCX_OUTPUT),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
            OpenAndRepair=True,
        )
        try:
            headings = []
            for index in range(1, check.Paragraphs.Count + 1):
                paragraph = check.Paragraphs.Item(index)
                style_name = str(paragraph.Style.NameLocal)
                if style_name in {"Titre 1", "Titre 2", "Titre 3"}:
                    text = clean_paragraph_text(paragraph).strip()
                    if text:
                        headings.append(text)

            result = {
                "docx": str(DOCX_OUTPUT),
                "pdf": str(PDF_OUTPUT),
                "pages": check.ComputeStatistics(WD_STATISTIC_PAGES),
                "words": check.ComputeStatistics(WD_STATISTIC_WORDS),
                "sections": check.Sections.Count,
                "tables": check.Tables.Count,
                "toc_count": check.TablesOfContents.Count,
                "headings": headings,
                "docx_bytes": DOCX_OUTPUT.stat().st_size,
                "pdf_bytes": PDF_OUTPUT.stat().st_size,
            }
        finally:
            check.Close(SaveChanges=False)
        return result
    finally:
        if document is not None:
            document.Close(SaveChanges=False)
        word.Quit()


if __name__ == "__main__":
    print(json.dumps(build_report(), ensure_ascii=False, indent=2))
