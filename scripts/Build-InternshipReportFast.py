from __future__ import annotations

import importlib.util
import json
import shutil
from pathlib import Path

import win32com.client
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "00-project" / "RAPPORT_STAGE_DISPLAY_CONTROL.docx"
OUT = ROOT / "docs" / "04-report" / "Rapport_PFE_Display_Control_NEGAGZA_Nizar.docx"
PDF = OUT.with_suffix(".pdf")


def load_repair_helpers():
    path = Path(__file__).with_name("Build-InternshipReport.py")
    spec = importlib.util.spec_from_file_location("report_repair", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def find_body_paragraph(document, exact: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise RuntimeError(f"Paragraphe introuvable : {exact}")


def replace_all(document, old: str, new: str, required: bool = True):
    count = 0
    for paragraph in all_paragraphs(document):
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new)
            count += 1
    if required and not count:
        raise RuntimeError(f"Texte introuvable : {old}")


def replace_paragraph(document, old: str, new: str):
    for paragraph in all_paragraphs(document):
        if paragraph.text.strip() == old:
            paragraph.text = new
            return paragraph
    raise RuntimeError(f"Paragraphe introuvable : {old}")


def direct_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def move_block(document, start: str, end: str, before: str):
    body = document._element.body
    children = list(body)

    def index_of(text: str) -> int:
        for index, element in enumerate(list(body)):
            if element.tag == qn("w:p") and direct_text(element) == text:
                return index
        raise RuntimeError(f"Bloc introuvable : {text}")

    start_index = index_of(start)
    end_index = index_of(end)
    block = children[start_index:end_index]
    for element in block:
        body.remove(element)
    destination = index_of(before)
    for offset, element in enumerate(block):
        body.insert(destination + offset, element)


def set_style(document, text: str, style: str):
    find_body_paragraph(document, text).style = document.styles[style]


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def main():
    helpers = load_repair_helpers()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUT)
    helpers.ensure_content_types(OUT)
    helpers.rebuild_as_valid_docx(OUT)
    document = Document(OUT)

    replacements = {
        "RAPPORT DE STAGE": "RAPPORT DE STAGE / PROJET DE FIN D’ÉTUDES",
        "[Nom et prénom de lʼétudiant]": "M. NEGAGZA Nizar",
        "Période : août – septembre 2026  |  Durée : six semaines": "Période : du 3 août au 11 septembre 2026  |  Durée : six semaines",
        "Document préparé le 27 août 2026": "Soutenance : [date à compléter]  |  Année universitaire : 2025–2026",
        "Information à compléter avant dépôt final": "Repères institutionnels vérifiés",
        "Ajouter la présentation officielle dʼElastomer Solutions, son secteur dʼactivité exact, son organisation et le service dʼaccueil après validation par lʼencadrant. Aucune donnée non confirmée nʼa été inventée dans cette version.": "Elastomer Solutions Maroc SARL — industrie automobile — Zone Franche dʼExportation, îlot 64, lot 3, 90100 Tanger. Source : site officiel, consulté le 31 août 2026. Le service dʼaccueil reste à confirmer avec lʼencadrant professionnel.",
        "Chapitre 3 : Réalisation de la solution": "Chapitre 3 : Réalisation, tests et résultats",
        "3.12 Conclusion": "3.12 Synthèse de la réalisation",
        "Chapitre 4 : Tests, résultats et bilan": "3.13 Stratégie de tests et résultats",
        "4.1 Introduction": "3.13.1 Introduction",
        "4.2 Stratégie de vérification": "3.13.2 Stratégie de vérification",
        "4.3 Test du Raspberry Pi virtuel": "3.13.3 Test du Raspberry Pi virtuel",
        "4.4 Résultats observés": "3.13.4 Résultats observés",
        "4.5 Anomalies et limites découvertes": "3.13.5 Anomalies corrigées et limites ouvertes",
        "4.6 Validation matérielle restant à réaliser": "3.13.6 Validation matérielle restant à réaliser",
        "4.7 Apports du stage": "3.13.7 Apports du stage",
        "4.8 Conclusion": "3.14 Conclusion",
        "Tableau 13 — Résultats du test virtuel du 27 août 2026": "Tableau 13 — Résultats de la campagne locale du 30 août 2026",
        "État au 27 août 2026": "État au 31 août 2026",
        "Tests agent actuels": "Campagne logicielle du 30 août 2026",
        "12/12 réussis": "102/102 .NET ; 8/8 React ; 5/5 navigateur",
        "Suite exécutée après le test virtuel": "Compilation, intégration, interfaces et accessibilité automatisée",
        "Générés par ce document": "Générés le 31 août 2026",
    }
    for old, new in replacements.items():
        replace_all(document, old, new)

    table = document.tables[0]
    values = [
        ("Information", "Détail"),
        ("Présenté par", "M. NEGAGZA Nizar"),
        ("Diplôme / filière / établissement", "Ingénieur d’État en Génie [à compléter] — ENSI Tanger"),
        ("Encadrant professionnel", "[À compléter]"),
        ("Encadrant pédagogique", "[À compléter]"),
        ("Année universitaire", "2025–2026"),
    ]
    for row, (left, right) in zip(table.rows, values):
        row.cells[0].text = left
        row.cells[1].text = right

    thanks = find_body_paragraph(document, "Remerciements")
    dedication = document.add_paragraph("Dédicace", style="Heading 1")
    thanks._p.addprevious(dedication._p)
    first = document.add_paragraph(
        "À mes parents et à ma famille, pour leur soutien constant, leur confiance et leurs encouragements tout au long de mon parcours."
    )
    dedication._p.addnext(first._p)
    second = document.add_paragraph(
        "À mes enseignants, à mes encadrants et à toutes les personnes qui ont contribué à ma formation et à la réussite de ce projet."
    )
    first._p.addnext(second._p)
    second.add_run().add_break(WD_BREAK.PAGE)

    move_block(document, "Table des matières", "Introduction générale", "Liste des acronymes")
    move_block(document, "Liste des acronymes", "Liste des figures", "Introduction générale")
    replace_all(document, "Liste des acronymes", "Liste des abréviations")

    replace_paragraph(
        document,
        "Le stage est réalisé au sein dʼElastomer Solutions à Tanger. Dans ce rapport, la présentation de lʼentreprise reste volontairement limitée aux informations confirmées par le stagiaire : son nom, son implantation à Tanger et son rôle dʼorganisme dʼaccueil. Les données institutionnelles détaillées, lʼorganigramme, les chiffres dʼactivité et les éléments confidentiels devront être ajoutés uniquement après validation par lʼentreprise.",
        "Le stage est réalisé au sein dʼElastomer Solutions Maroc SARL, implantée dans la Zone Franche dʼExportation de Tanger. Le groupe développe et fabrique des composants en caoutchouc, en thermoplastique et des pièces bi-matière caoutchouc-plastique destinés principalement à lʼindustrie automobile. Le site marocain a démarré lʼinjection de composants en caoutchouc en 2013 ; une nouvelle unité de production à Tanger est entrée en activité en 2016 afin de servir les marchés dʼEurope, dʼAfrique et des Amériques [1].",
    )

    outdated = [
        (
            "Le test a identifié un défaut du script Windows Start-VirtualDevice.ps1 : lorsquʼil utilise dotnet run depuis le dossier source, lʼAPI locale du lecteur fonctionne mais les fichiers React compilés sont recherchés dans un autre répertoire, ce qui provoque une réponse 404 sur la racine. Lʼexécution depuis le dossier compilé, équivalente au paquet publié sur le Pi, sert correctement lʼinterface. Le script de simulation doit être ajusté pour reproduire ce comportement sans manipulation manuelle.",
            "Le premier essai virtuel a révélé un défaut de chemin dans le lanceur Windows. Le lanceur et la préparation de lʼenvironnement ont ensuite été corrigés ; le parcours complet est désormais reproductible depuis le paquet compilé prévu pour le Raspberry Pi.",
        ),
        (
            "Un scénario dʼintégration complet a également rencontré une attente devenue obsolète après le renforcement de la politique de mot de passe. Le mot de passe de test contenait le mot « viewer », identique à la partie locale de lʼadresse électronique, et il a été correctement refusé comme information personnelle. Le comportement de sécurité est cohérent, mais la donnée du test doit être mise à jour avant de compter ce scénario comme vert dans la vérification actuelle.",
            "Un scénario dʼintégration a mis en évidence une donnée de test incompatible avec la politique de mot de passe renforcée. La fixture a été corrigée sans affaiblir la règle de sécurité, puis toute la campagne a été relancée.",
        ),
        (
            "Ces résultats illustrent lʼintérêt dʼun rapport honnête. Les preuves antérieures du dépôt indiquent une campagne complète réussie avant les derniers changements. Après ces changements, la suite agent passe et le parcours virtuel fonctionne, mais la campagne complète doit être relancée après correction de la fixture de test. Aucune affirmation de préparation à la production ne doit reposer uniquement sur le test virtuel.",
            "La vérification locale du 30 août 2026 totalise 102 tests .NET, 8 tests React et 5 scénarios Playwright/axe, tous réussis. La compilation Release se termine sans avertissement ni erreur. Ces preuves ne remplacent pas les essais sur Raspberry Pi physique ni la validation dʼune infrastructure de recette.",
        ),
    ]
    for old, new in outdated:
        replace_paragraph(document, old, new)

    set_style(document, "3.13 Stratégie de tests et résultats", "Heading 2")
    for text in replacements.values():
        if isinstance(text, str) and text.startswith("3.13."):
            set_style(document, text, "Heading 3")
    set_style(document, "3.14 Conclusion", "Heading 2")

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.page_break_before = True
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in document.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.footer_distance = Cm(1)
        footer = section.footer
        footer.paragraphs[0].clear()
        add_page_number(footer.paragraphs[0])

    document.core_properties.title = "Rapport de stage / PFE — Display Control"
    document.core_properties.author = "M. NEGAGZA Nizar"
    document.save(OUT)

    if PDF.exists():
        PDF.unlink()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(OUT), ReadOnly=True, AddToRecentFiles=False)
    try:
        doc.Fields.Update()
        pages = doc.ComputeStatistics(2)
        words = doc.ComputeStatistics(0)
        doc.ExportAsFixedFormat(str(PDF), 17)
    finally:
        doc.Close(False)
        word.Quit()

    print(json.dumps({"docx": str(OUT), "pdf": str(PDF), "pages": pages, "words": words}, ensure_ascii=False))


if __name__ == "__main__":
    main()
